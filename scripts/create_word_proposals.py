# -*- coding: utf-8 -*-
"""
Générateur du Document Word Officiel :
Propositions Méthodologiques Approfondies - Projet SIMACCUEIL
Haute École Charlemagne (HECh) / Haute École de Liège (HEL) / Lab-DRA
Chercheur Principal : Jérôme Foguenne
"""

import os
import sys
import shutil

if sys.stdout:
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

LOCAL_DOCX = r"C:\Users\TrendingPC\.gemini\antigravity\scratch\Eye-Tracking\docs\Propositions_Methodologiques_SIMACCUEIL.docx"
DRIVE_DOCX = r"C:\Google Drive\Prépas light\HECh\Projet recherche\Propositions Methodologiques - SIMACCUEIL.docx"

COLOR_NAVY = RGBColor(27, 54, 93)      # #1B365D
COLOR_TEAL = RGBColor(14, 134, 212)    # #0E86D4
COLOR_CHARCOAL = RGBColor(44, 62, 80)  # #2C3E50
COLOR_MUTED = RGBColor(100, 110, 120)  # #646E78

HEX_NAVY = "1B365D"
HEX_TEAL = "0E86D4"
HEX_LIGHT_BG = "F4F6F9"
HEX_ALT_ROW = "F9FAFC"
HEX_BORDER = "D1D5DB"

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)
        
        # Header
        hdr = s.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.text = "Lab-DRA | SIMACCUEIL — Trois Propositions Méthodologiques de Recherche"
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_hdr.style.font.name = "Calibri"
        p_hdr.style.font.size = Pt(8.5)
        p_hdr.style.font.color.rgb = COLOR_MUTED
        
        # Footer
        ftr = s.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.text = "HECh / HEL — Jérôme Foguenne (Chercheur Principal)  |  Rapport Méthodologique Officiel"
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ftr.style.font.name = "Calibri"
        p_ftr.style.font.size = Pt(8.5)
        p_ftr.style.font.color.rgb = COLOR_MUTED

def style_cell(cell, bg_color=None, top_pad=100, bot_pad=100, left_pad=140, right_pad=140, border_left=None):
    tcPr = cell._tc.get_or_add_tcPr()
    if bg_color:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{bg_color}"/>')
        tcPr.append(shd)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top_pad}" w:type="dxa"/><w:bottom w:w="{bot_pad}" w:type="dxa"/><w:left w:w="{left_pad}" w:type="dxa"/><w:right w:w="{right_pad}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)
    if border_left:
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="{border_left.get("sz", 24)}" w:space="0" w:color="{border_left.get("color", HEX_NAVY)}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)

def build_proposals_docx():
    doc = Document()
    add_header_footer(doc)
    
    # Page de titre
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(6)
    r_inst = p_inst.add_run("HAUTE ÉCOLE CHARLEMAGNE (HECh) — HAUTE ÉCOLE DE LIÈGE (HEL)\nLABORATOIRE DE RECHERCHE LAB-DRA")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_TEAL

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("DOCUMENT D'ORIENTATION MÉTHODOLOGIQUE")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Trois Propositions Méthodologiques Détaillées pour le Projet SIMACCUEIL\nOculométrie Mobile (Pupil Labs Neon), Didactique Professionnelle & Accueil Hôtelier (Upselling & Réclamation)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12.5)
    r_sub.font.color.rgb = COLOR_CHARCOAL

    # Table cartouche métadonnées
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Chercheur Principal", "Jérôme Foguenne (HECh / HEL)"),
        ("Ancrage Théorique", "État de l'art consolidé (22 références académiques majeures)"),
        ("Dispositif Expérimental", "Lunettes d'Eye Tracking Mobile Pupil Labs Neon (200 Hz) & Simulateur PMS"),
        ("Finalité Stratégique", "Dépôt projets de recherche appliquée, financements compétitifs & innovation didactique")
    ]
    for r_idx, (k, v) in enumerate(meta_data):
        c0, c1 = meta_table.rows[r_idx].cells
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        style_cell(c0, bg_color=HEX_LIGHT_BG, top_pad=80, bot_pad=80, left_pad=120, right_pad=120)
        style_cell(c1, bg_color="FFFFFF", top_pad=80, bot_pad=80, left_pad=120, right_pad=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_NAVY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_CHARCOAL

    doc.add_page_break()

    # Introduction
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("1. Cadrage Épistémologique & Matrice Comparée des Propositions")
    r1.font.name = "Calibri"
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_NAVY

    p_in = doc.add_paragraph()
    p_in.paragraph_format.line_spacing = 1.15
    p_in.paragraph_format.space_after = Pt(10)
    p_in.add_run(
        "Le projet de recherche appliquée SIMACCUEIL étudie la relation de service en réception hôtelière à travers deux situations professionnelles "
        "critiques : l'upselling commercial (surclassement en chambre supérieure) et la gestion bienveillante d'une réclamation client. "
        "Pour répondre aux exigences de scientificité du Lab-DRA et aux finalités didactiques de l'HECh/HEL, trois propositions méthodologiques "
        "complètes et articulées à l'état de l'art consolidé (22 références) sont formulées ci-dessous."
    )

    # Tableau comparatif
    cmp_table = doc.add_table(rows=6, cols=4)
    cmp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Critère", "Proposition 1 : Comparatif", "Proposition 2 : EMME", "Proposition 3 : Dyadique"]
    for c_idx, h_text in enumerate(headers):
        cell = cmp_table.rows[0].cells[c_idx]
        style_cell(cell, bg_color=HEX_NAVY, top_pad=100, bot_pad=100, left_pad=100, right_pad=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Objectif Central", "Modéliser les marqueurs visuels de l'expertise en réception.", "Mesurer l'impact pédagogique d'un entraînement 'Spotlight'.", "Décortiquer la synchronie sociale et le vécu en interaction."),
        ("Design Expérimental", "Quasi-expérimental transversal (Novices vs Intermédiaires vs Experts).", "RCT interventionnel pré-test / post-test randomisé.", "Mixte séquentiel (Dual Eye Tracking + Autoconfrontation)."),
        ("Matériel Oculométrique", "1 paire de lunettes Neon (réceptionniste).", "1 paire de lunettes Neon (étudiant pré/post-tests).", "2 paires de lunettes Neon synchronisées (réceptionniste + client)."),
        ("Références Clés", "Beckmann (2010), Gegenfurtner (2011), Wolff (2016), Goodwin (1994), Cortina (2015), Vickers (2016).", "Jarodzka (2012), Seppänen (2012), Van Gog (2009), Elling (2012), Hooge (2024).", "Rogers (2018), Wohltjen (2021), Theureau (2016), Tickle-Degnen (1990), Crandall (2006)."),
        ("Livrable HECh", "Référentiel métrologique des compétences visuelles pour évaluation.", "Modules vidéo d'auto-formation par eye tracking prêts à l'emploi.", "Modèle d'ingénierie relationnelle pour la formation supérieure.")
    ]
    for r_idx, r_data in enumerate(rows_data):
        row = cmp_table.rows[r_idx + 1]
        bg = HEX_ALT_ROW if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            style_cell(cell, bg_color=bg, top_pad=60, bot_pad=60, left_pad=80, right_pad=80)
            p = cell.paragraphs[0]
            if c_idx == 0:
                p.add_run(val).font.bold = True
                p.runs[0].font.size = Pt(8.5)
            else:
                p.add_run(val).font.size = Pt(8.5)

    doc.add_page_break()

    # PROPOSITION 1
    h_p1 = doc.add_paragraph()
    h_p1.paragraph_format.space_before = Pt(12)
    h_p1.paragraph_format.space_after = Pt(6)
    r_p1 = h_p1.add_run("2. PROPOSITION 1 : Approche Comparativo-Développementale (Experts vs Novices)")
    r_p1.font.name = "Calibri"
    r_p1.font.size = Pt(15)
    r_p1.font.bold = True
    r_p1.font.color.rgb = COLOR_NAVY

    doc.add_paragraph("Titre de recherche : « La Vision Professionnelle en Action : Étude comparative transversale de l'allocation attentionnelle, de la charge cognitive et de la performance relationnelle en accueil hôtelier instrumenté »").runs[0].font.italic = True

    # Section 1.1 Justification
    p_j1 = doc.add_paragraph()
    p_j1.paragraph_format.space_before = Pt(8)
    p_j1.add_run("2.1. Justification détaillée des choix scientifiques (Ancrage État de l'Art)").font.bold = True
    p_j1.runs[0].font.color.rgb = COLOR_TEAL

    p_j1_txt = doc.add_paragraph()
    p_j1_txt.paragraph_format.line_spacing = 1.15
    p_j1_txt.add_run(
        "Cette première méthodologie opérationnalise la méta-analyse de Gegenfurtner, Lehtinen et Säljö (2011) sur les lois universelles du regard expert : "
        "réduction de l'information (abandon des distracteurs), durées de fixation plus courtes et traitement parafovéal étendu. "
        "Elle intègre le cadre de complexité de Beckmann (2010) : la manipulation du PMS génère une forte charge cognitive extrinsèque chez l'étudiant novice. "
        "L'automatisation des routines informatiques chez le professionnel expérimenté libère des ressources en mémoire de travail, lui permettant d'orienter son attention vers le client. "
        "Cette proposition teste également le modèle de la vision professionnelle de Goodwin (1994) et de Wolff et al. (2016), la convergence prédictive de Cortina et al. (2015) "
        "et la stabilisation du Quiet Eye de Vickers (2016) avant l'annonce du tarif d'upselling (Denizci Guillet, 2020 ; Ruki, 2023)."
    )

    # Section 1.2 Plan
    p_pl1 = doc.add_paragraph()
    p_pl1.add_run("2.2. Plan du projet de recherche (Phasage sur 12 mois)").font.bold = True
    p_pl1.runs[0].font.color.rgb = COLOR_TEAL
    
    plan1_steps = [
        ("Mois 1-2 (Cadrage & Pilote)", "Scénarisation des micro-situations (Upselling & Réclamation), paramétrage du PMS expérimental et calibration des AOI selon Hooge et al. (2024)."),
        ("Mois 3-6 (Recueil Expérimental)", "Passation de 60 participants sous lunettes Neon 200 Hz (20 Novices B1, 20 Intermédiaires B3, 20 Professionnels 4*-5*)."),
        ("Mois 7-8 (Traitement & IA)", "Extraction automatisée des fixations via le pipeline de vision par ordinateur de De Beugher et al. (2014) et double codage en aveugle de la qualité d'accueil."),
        ("Mois 9-12 (Analyses & Valorisation)", "Analyses ANOVA/MANOVA, modélisation de l'expertise, rédaction du référentiel pour l'HECh et soumission d'article scientifique.")
    ]
    for s_title, s_desc in plan1_steps:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.left_indent = Inches(0.2)
        p_s.paragraph_format.space_after = Pt(2)
        r_st = p_s.add_run(f"• {s_title} : ")
        r_st.font.bold = True
        p_s.add_run(s_desc)

    # Section 1.3 Objectifs & Grille
    p_g1 = doc.add_paragraph()
    p_g1.paragraph_format.space_before = Pt(8)
    p_g1.add_run("2.3. Grille des indicateurs opérationnels (Proposition 1)").font.bold = True
    p_g1.runs[0].font.color.rgb = COLOR_TEAL

    t_g1 = doc.add_table(rows=6, cols=4)
    t_g1.alignment = WD_TABLE_ALIGNMENT.CENTER
    g1_headers = ["Indicateur", "Métrique & Calcul", "Outil de Recueil", "Justification État de l'Art"]
    for c_idx, h_text in enumerate(g1_headers):
        cell = t_g1.rows[0].cells[c_idx]
        style_cell(cell, bg_color=HEX_NAVY, top_pad=80, bot_pad=80, left_pad=80, right_pad=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    g1_rows = [
        ("Dwell Time PMS (%)", "Pourcentage du temps total passé à fixer l'interface logicielle.", "Pupil Labs Neon (Surface AOI)", "Beckmann (2010) : Charge extrinsèque liée au manque d'automatisation."),
        ("TTFF - Champ Tarifaire", "Latence (ms) jusqu'à la 1re fixation sur le champ tarifaire / surclassement.", "Neon (AOI Statique PMS)", "Gegenfurtner et al. (2011) : Accès parafovéal plus rapide chez l'expert."),
        ("Ratio Regard Client / PMS", "Temps de fixation Visage Client / Temps de fixation Écran PMS.", "Calcul automatique coordonnées", "Wolff et al. (2016) : Mesure de la vision panoramique vs cécité écran."),
        ("Quiet Eye Duration (QED)", "Durée (ms) de la dernière fixation stable avant la prise de parole critique.", "Microphone + Neon", "Vickers (2016) : Programmation cognitive et motrice optimale."),
        ("Score de Rapport Non Verbal", "Score composite sur 7 points (attention mutuelle, positivité, coordination).", "Grille CLASS (Cortina 2015)", "Tickle-Degnen (1990) & Cortina (2015) : Validation par convergence.")
    ]
    for r_idx, r_data in enumerate(g1_rows):
        row = t_g1.rows[r_idx + 1]
        bg = HEX_ALT_ROW if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            style_cell(cell, bg_color=bg, top_pad=50, bot_pad=50, left_pad=60, right_pad=60)
            cell.paragraphs[0].add_run(val).font.size = Pt(8)

    doc.add_page_break()

    # PROPOSITION 2
    h_p2 = doc.add_paragraph()
    h_p2.paragraph_format.space_before = Pt(12)
    h_p2.paragraph_format.space_after = Pt(6)
    r_p2 = h_p2.add_run("3. PROPOSITION 2 : Approche Didactique & Vicariante (Interventionnelle par EMME)")
    r_p2.font.name = "Calibri"
    r_p2.font.size = Pt(15)
    r_p2.font.bold = True
    r_p2.font.color.rgb = COLOR_NAVY

    doc.add_paragraph("Titre de recherche : « Ingénierie Pédagogique par Modélisation du Regard (EMME) : Évaluation de l'impact d'un guidage attentionnel 'Spotlight' sur l'apprentissage des compétences d'accueil hôtelier »").runs[0].font.italic = True

    # 2.1 Justification
    p_j2 = doc.add_paragraph()
    p_j2.paragraph_format.space_before = Pt(8)
    p_j2.add_run("3.1. Justification détaillée des choix scientifiques (Ancrage État de l'Art)").font.bold = True
    p_j2.runs[0].font.color.rgb = COLOR_TEAL

    p_j2_txt = doc.add_paragraph()
    p_j2_txt.paragraph_format.line_spacing = 1.15
    p_j2_txt.add_run(
        "Cette méthodologie interventionnelle s'appuie directement sur les travaux de Jarodzka et al. (2012) et de Seppänen & Gegenfurtner (2012). "
        "Elle répond à la mise en garde de Van Gog et al. (2009) en évitant les cercles oculaires distrayants au profit de la technique du 'Spotlight' "
        "(zone fixée nette, environnement périphérique estompé). L'apprentissage par observation des saccades expertes permet aux étudiants "
        "d'assimiler l'organisation visuelle du check-in sans surcharge cognitive. L'analyse des silences d'Elling, Lentz et de Jong (2012) est mobilisée "
        "pour mesurer l'allègement de la charge cognitive lors du post-test."
    )

    # 2.2 Plan
    p_pl2 = doc.add_paragraph()
    p_pl2.add_run("3.2. Plan du projet de recherche (Phasage sur 10 mois)").font.bold = True
    p_pl2.runs[0].font.color.rgb = COLOR_TEAL
    
    plan2_steps = [
        ("Mois 1-3 (Production Didactique)", "Tournage des 4 scénarios experts sous lunettes Neon et développement algorithmique du filtre vidéo Spotlight (FFmpeg/Python)."),
        ("Mois 4-6 (Essai Randomisé RCT)", "Passation du pré-test T0 (50 étudiants sous Neon), intervention didactique (Bras EMME Spotlight vs Bras Contrôle classique), puis post-test T1."),
        ("Mois 7-8 (Traitement & Métriques)", "Calcul automatisé des deltas de progression Delta(T1 - T0) et analyse de la réduction des silences non résolus face au client."),
        ("Mois 9-10 (Transfert Curriculaire)", "Mesure de la charge mentale subjective (échelle de Paas) et intégration pérenne des capsules vidéo dans Moodle HECh/HEL.")
    ]
    for s_title, s_desc in plan2_steps:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.left_indent = Inches(0.2)
        p_s.paragraph_format.space_after = Pt(2)
        r_st = p_s.add_run(f"• {s_title} : ")
        r_st.font.bold = True
        p_s.add_run(s_desc)

    # 2.3 Grille
    p_g2 = doc.add_paragraph()
    p_g2.paragraph_format.space_before = Pt(8)
    p_g2.add_run("3.3. Grille des indicateurs opérationnels (Proposition 2)").font.bold = True
    p_g2.runs[0].font.color.rgb = COLOR_TEAL

    t_g2 = doc.add_table(rows=6, cols=4)
    t_g2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h_text in enumerate(g1_headers):
        cell = t_g2.rows[0].cells[c_idx]
        style_cell(cell, bg_color=HEX_NAVY, top_pad=80, bot_pad=80, left_pad=80, right_pad=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    g2_rows = [
        ("Gain Attentionnel Client", "Progression du % temps regard dirigé vers le client : Delta = T1 - T0.", "Neon (AOI Visage Client)", "Jarodzka et al. (2012) : Transfert réussi du balayage visuel."),
        ("Réduction Errance PMS", "Diminution du temps d'errance sur les menus PMS : Delta = T0 - T1.", "Neon (AOI Interface PMS)", "Seppänen & Gegenfurtner (2012) : Élimination des fixations parasites."),
        ("Silences Non Résolus", "Durée cumulée (s) des pauses verbales > 2s avec regard bloqué sur l'écran.", "Audio + Oculométrie", "Elling et al. (2012) : Révélateur direct de saturation cognitive."),
        ("Score d'Effort Mental", "Évaluation psychométrique de la charge ressentie sur l'échelle de Paas (1 à 9).", "Questionnaire post-tâche", "Beckmann (2010) : Diminution de la charge extrinsèque."),
        ("Qualité de l'Upselling", "Grille critériée sur 20 points (conformité tarifaire, élocution, persuasion).", "Grille d'examen HECh", "Ruki (2023) : Maîtrise des produits et aisance commerciale.")
    ]
    for r_idx, r_data in enumerate(g2_rows):
        row = t_g2.rows[r_idx + 1]
        bg = HEX_ALT_ROW if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            style_cell(cell, bg_color=bg, top_pad=50, bot_pad=50, left_pad=60, right_pad=60)
            cell.paragraphs[0].add_run(val).font.size = Pt(8)

    doc.add_page_break()

    # PROPOSITION 3
    h_p3 = doc.add_paragraph()
    h_p3.paragraph_format.space_before = Pt(12)
    h_p3.paragraph_format.space_after = Pt(6)
    r_p3 = h_p3.add_run("4. PROPOSITION 3 : Approche Énactive & Interactionnelle Dyadique (Dual Eye Tracking & Cours d'Action)")
    r_p3.font.name = "Calibri"
    r_p3.font.size = Pt(15)
    r_p3.font.bold = True
    r_p3.font.color.rgb = COLOR_NAVY

    doc.add_paragraph("Titre de recherche : « Énaction et Synchronie Sociale : Analyse dyadique de l'interaction d'accueil hôtelier par double oculométrie mobile synchronisée et entretiens d'autoconfrontation du cours d'action »").runs[0].font.italic = True

    # 3.1 Justification
    p_j3 = doc.add_paragraph()
    p_j3.paragraph_format.space_before = Pt(8)
    p_j3.add_run("4.1. Justification détaillée des choix scientifiques (Ancrage État de l'Art)").font.bold = True
    p_j3.runs[0].font.color.rgb = COLOR_TEAL

    p_j3_txt = doc.add_paragraph()
    p_j3_txt.paragraph_format.line_spacing = 1.15
    p_j3_txt.add_run(
        "Cette proposition de pointe dépasse l'analyse unilatérale en équipant simultanément le réceptionniste et le client de lunettes Neon. "
        "Elle s'ancre dans les découvertes de Rogers et al. (2018) sur les dynamiques du regard social face-à-face et de Wohltjen & Wheatley (2021) "
        "sur les cycles d'attention partagée et de désynchronisation salutaire. Sur le plan qualitatif, elle mobilise le cadre du cours d'action "
        "de Theureau (Durand, 2016) et la Cognitive Task Analysis de Crandall, Klein et Hoffman (2006). L'entretien d'autoconfrontation "
        "immédiat permet à l'étudiant d'expliciter son vécu subjectif, ses dilemmes et ses intentions lors des moments de tension."
    )

    # 3.2 Plan
    p_pl3 = doc.add_paragraph()
    p_pl3.add_run("4.2. Plan du projet de recherche (Phasage sur 12 mois)").font.bold = True
    p_pl3.runs[0].font.color.rgb = COLOR_TEAL
    
    plan3_steps = [
        ("Mois 1-2 (R&D Métrologique)", "Configuration de la synchronisation réseau local NTP inter-lunettes Neon (< 5 ms) et scénographie du comptoir avec tablette orientable (Otterbring et al., 2014)."),
        ("Mois 3-5 (Collecte Dyadique)", "20 passations dyadiques complètes (2 paires Neon) immédiatement suivies d'entretiens d'autoconfrontation de 45 minutes par étudiant (Theureau, 2016)."),
        ("Mois 6-8 (Modélisation Mixte)", "Calcul de la synchronie temporelle par récurrence croisée (CRQA) et découpage sémiotique du verbatim en unités d'action-signe."),
        ("Mois 9-12 (Synthèse & Valorisation)", "Croisement qualitatif/quantitatif des dilemmes vécus et rédaction d'une monographie sur l'intelligence relationnelle en hospitalité.")
    ]
    for s_title, s_desc in plan3_steps:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.left_indent = Inches(0.2)
        p_s.paragraph_format.space_after = Pt(2)
        r_st = p_s.add_run(f"• {s_title} : ")
        r_st.font.bold = True
        p_s.add_run(s_desc)

    # 3.3 Grille
    p_g3 = doc.add_paragraph()
    p_g3.paragraph_format.space_before = Pt(8)
    p_g3.add_run("4.3. Grille des indicateurs opérationnels (Proposition 3)").font.bold = True
    p_g3.runs[0].font.color.rgb = COLOR_TEAL

    t_g3 = doc.add_table(rows=6, cols=4)
    t_g3.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c_idx, h_text in enumerate(g1_headers):
        cell = t_g3.rows[0].cells[c_idx]
        style_cell(cell, bg_color=HEX_NAVY, top_pad=80, bot_pad=80, left_pad=80, right_pad=80)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    g3_rows = [
        ("Taux de Contact Visuel Direct", "% de temps où les vecteurs de regard des deux acteurs se croisent mutuellement.", "Calcul intersection 2x Neon", "Rogers et al. (2018) : Discontinuité réelle du contact visuel."),
        ("Attention Conjointe Tablette", "Durée cumulée où le réceptionniste et le client fixent la tablette ensemble (+/- 200 ms).", "Détection conjointe AOI", "Wohltjen & Wheatley (2021) & Otterbring (2014) : Co-attention d'upselling."),
        ("Contagion Émotionnelle", "Délai d'alignement temporel entre le sourire de l'employé et l'apaisement facial du client.", "Codage FACS vidéo Neon", "Hennig-Thurau et al. (2006) : Authenticité émotionnelle réelle."),
        ("Unités d'Action-Signe (UAS)", "Décomposition de l'activité vécue : Engagement, Prise d'indice, Interprétation, Action.", "Verbatim autoconfrontation", "Theureau / Durand (2016) : Reconstitution du sens pour l'acteur."),
        ("Dilemmes Décisionnels Vécus", "Moments de tension explicités par l'étudiant face aux contraintes du logiciel et du client.", "Questions CDM (Crandall)", "Crandall, Klein & Hoffman (2006) : Explicitation des savoirs tacites.")
    ]
    for r_idx, r_data in enumerate(g3_rows):
        row = t_g3.rows[r_idx + 1]
        bg = HEX_ALT_ROW if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            style_cell(cell, bg_color=bg, top_pad=50, bot_pad=50, left_pad=60, right_pad=60)
            cell.paragraphs[0].add_run(val).font.size = Pt(8)

    doc.add_page_break()

    # Recommandation Finale
    h_rec = doc.add_paragraph()
    h_rec.paragraph_format.space_before = Pt(12)
    h_rec.paragraph_format.space_after = Pt(6)
    r_rec = h_rec.add_run("5. Recommandation Stratégique pour le Chercheur Principal")
    r_rec.font.name = "Calibri"
    r_rec.font.size = Pt(15)
    r_rec.font.bold = True
    r_rec.font.color.rgb = COLOR_NAVY

    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.line_spacing = 1.15
    p_rec.add_run(
        "Pour maximiser la rentabilité scientifique et l'impact pédagogique du projet SIMACCUEIL, il est préconisé un déploiement séquentiel en entonnoir :\n\n"
        "1. Année 1 (Proposition 1 - Comparatif Experts/Novices) : Rodage du laboratoire, calibration fine des AOI selon Hooge et al. (2024), "
        "publication internationale de référence sur la 'Vision Professionnelle' en hospitalité.\n"
        "2. Année 2 (Proposition 2 - Intervention Didactique EMME) : Traduction des marqueurs d'expertise en capsules pédagogiques 'Spotlight' "
        "évaluées par essai randomisé auprès des étudiants de l'HECh/HEL.\n"
        "3. Année 3 / Thèse (Proposition 3 - Dyadique Dual Eye Tracking) : Exploration en profondeur de la synchronie sociale et du cours d'action "
        "pour asseoir la réputation internationale du Lab-DRA dans le domaine des interactions de service avancées."
    )

    # Sauvegarde
    os.makedirs(os.path.dirname(LOCAL_DOCX), exist_ok=True)
    doc.save(LOCAL_DOCX)
    print(f"✅ Document Word généré avec succès dans : {LOCAL_DOCX}")

    try:
        shutil.copy2(LOCAL_DOCX, DRIVE_DOCX)
        print(f"✅ Document copié avec succès sur Google Drive : {DRIVE_DOCX}")
    except PermissionError:
        alt_drive = DRIVE_DOCX.replace(".docx", " (Mis a jour).docx")
        shutil.copy2(LOCAL_DOCX, alt_drive)
        print(f"⚠️ Fichier verrouillé par Word, copié sous : {alt_drive}")

if __name__ == "__main__":
    build_proposals_docx()
