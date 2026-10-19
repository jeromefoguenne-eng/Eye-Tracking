import docx
import os
import glob

# Find path reliably
drive_folder = r"C:\Google Drive"
matches = []
for root, dirs, files in os.walk(drive_folder):
    for f in files:
        if "Constats" in f:
            matches.append(os.path.join(root, f))

print("Found files:", matches)
if matches:
    target_path = matches[0]
    print(f"Opening: {target_path}")
    doc = docx.Document(target_path)
    
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"P{i}: {p.text}")
            
    print(f"\nTotal tables: {len(doc.tables)}")
    for t_idx, t in enumerate(doc.tables):
        print(f"\n--- Table {t_idx} ---")
        for r_idx, row in enumerate(t.rows):
            cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
            print(f"R{r_idx}: {cells}")
